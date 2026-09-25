"""
AR Subledger to GL 1200 Tie-Out (monthly)

Fictional example for Fernhill Supply Co. All data is synthetic.

What it does:
    1. Reads the AR subledger (open invoices) and the GL trial balance
       for one reporting period.
    2. Excludes intercompany customers.
    3. Compares the adjusted subledger total to GL account 1200.
    4. Writes an Excel workbook with Recon, Exceptions and Control tabs.

How to run (see HUMAN_INSTRUCTIONS.md for the full steps):
    pip install -r requirements.txt
    python sample_code.py

Reminder: the workflow running without an error does not mean the
accounting result is correct. A person still reviews the output.
"""

import getpass
import re
import subprocess
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# PERIOD UPDATE:
# Update REPORT_MONTH and REPORT_YEAR before each monthly run.
# Input files must use YYYYMM in the filename (for example
# ar_subledger_202609.csv). The script refuses to run if the files in
# data/ are for a different month, so last month's files cannot be
# reconciled by accident.
# ---------------------------------------------------------------------------
REPORT_MONTH = 9
REPORT_YEAR = 2026

WORKFLOW_NAME = "AR Subledger to GL 1200 Tie-Out"
WORKFLOW_VERSION = "1.0.0"  # Bump this and add a CHANGELOG line when the logic changes.

# Where the official copy of this workflow lives. The Control tab prints this
# so a reviewer (or auditor) can find the workflow from the output alone.
# Change it if your team keeps the workflow somewhere else.
REPOSITORY_LOCATION = (
    "https://github.com/PythonMuse/pythonmuse-builder-to-production/"
    "tree/main/examples/monthly_reconciliation_example"
)

# BUSINESS RULE:
# GL account 1200 is Accounts Receivable - Trade. Intercompany balances are
# booked to 1210, which is why intercompany customers are excluded below.
GL_AR_ACCOUNT = "1200"

# VALIDATION:
# Final AR balance must agree to GL account 1200 within $0.01.
# Anything larger is written to the Exceptions tab for manual review.
# $0.01 allows for rounding only. It is not a materiality threshold.
TOLERANCE = 0.01

HERE = Path(__file__).resolve().parent
DATA_DIR = HERE / "data"
OUTPUT_DIR = HERE / "output"


def fail(message):
    """Stop the run with a plain-English message instead of a stack trace."""
    print("\nSTOPPED: " + message + "\n")
    sys.exit(1)


def find_input(prefix, period):
    """Return the input file for this period, or stop with a clear reason."""
    expected = DATA_DIR / f"{prefix}_{period}.csv"
    if expected.exists():
        return expected

    candidates = sorted(DATA_DIR.glob(f"{prefix}*.csv"))
    if not candidates:
        fail(f"No '{prefix}' file found in {DATA_DIR}. Expected {expected.name}.")

    for path in candidates:
        # VALIDATION:
        # Every input filename must carry its period as YYYYMM. A file called
        # "ar_subledger_final_v2.csv" tells nobody which month it belongs to.
        if not re.search(r"(?<!\d)\d{6}(?!\d)", path.stem):
            fail(
                f"'{path.name}' does not contain a YYYYMM period in its name. "
                f"Rename it, for example to {expected.name}."
            )

    found = ", ".join(p.name for p in candidates)
    fail(
        f"No '{prefix}' file for period {period}. Found: {found}. "
        "Either the new month's file has not been saved to data/ yet, "
        "or REPORT_MONTH / REPORT_YEAR still need updating."
    )


def file_period(path):
    """The YYYYMM period printed in an input filename."""
    return re.search(r"(?<!\d)\d{6}(?!\d)", path.stem).group()


def git_commit():
    """Short commit hash of this workflow, with graceful fallbacks."""
    try:
        commit = subprocess.run(
            ["git", "rev-parse", "--short", "HEAD"],
            cwd=HERE, capture_output=True, text=True, timeout=10,
        )
    except (OSError, subprocess.SubprocessError):
        return "unknown (git not available)"
    if commit.returncode != 0:
        # Either not a git repository, or a repository with no commits yet.
        inside = subprocess.run(
            ["git", "rev-parse", "--is-inside-work-tree"],
            cwd=HERE, capture_output=True, text=True, timeout=10,
        )
        return "uncommitted" if inside.returncode == 0 else "unknown (not a git repository)"

    sha = commit.stdout.strip()
    # Flag local edits so nobody mistakes a modified copy for the approved one.
    status = subprocess.run(
        ["git", "status", "--porcelain", "--", "."],
        cwd=HERE, capture_output=True, text=True, timeout=10,
    )
    changed = [
        line for line in status.stdout.splitlines()
        if "output/" not in line  # The generated workbook is not a code change.
    ]
    return f"{sha} (uncommitted changes present)" if changed else sha


def load_subledger(path):
    # Read Excel file
    # ^ DELIBERATELY POOR COMMENT (kept on purpose as a teaching example).
    #   It is wrong (this reads a CSV) and, even if it were right, it only
    #   repeats what the code already says. It tells the next accountant
    #   nothing about WHY, or what to change. Compare it with the
    #   BUSINESS RULE / PERIOD UPDATE / VALIDATION comments in this file.
    df = pd.read_csv(path, dtype={"customer_id": str, "intercompany_flag": str})

    required = {"customer_id", "customer_name", "invoice_number",
                "open_balance", "intercompany_flag"}
    missing = required - set(df.columns)
    if missing:
        fail(f"{path.name} is missing column(s): {', '.join(sorted(missing))}. "
             "Has the AR aging export layout changed?")

    # VALIDATION:
    # The intercompany flag must be Y or N. A blank or unexpected value
    # stops the run rather than guessing, because a guess could silently
    # include or exclude a customer.
    df["intercompany_flag"] = df["intercompany_flag"].fillna("").str.strip().str.upper()
    bad = df[~df["intercompany_flag"].isin(["Y", "N"])]
    if not bad.empty:
        rows = ", ".join(bad["invoice_number"].astype(str))
        fail(f"intercompany_flag must be Y or N. Check invoice(s): {rows}.")
    return df


def load_gl_balance(path, period):
    gl = pd.read_csv(path, dtype={"account": str, "period": str})

    # VALIDATION:
    # The period inside the file must match the period in its filename.
    # This catches a file that was renamed without being re-exported.
    periods = set(gl["period"].str.strip())
    if periods != {period}:
        fail(f"{path.name} contains period(s) {sorted(periods)}, expected only {period}.")

    rows = gl[gl["account"].str.strip() == GL_AR_ACCOUNT]
    if len(rows) != 1:
        fail(f"Expected exactly one row for GL account {GL_AR_ACCOUNT} in {path.name}, "
             f"found {len(rows)}.")
    return round(float(rows["balance"].iloc[0]), 2)


def format_sheet(ws, money_cols=()):
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col_idx, column in enumerate(ws.columns, start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width + 2, 12), 80)
        if col_idx in money_cols:
            for c in column[1:]:
                if isinstance(c.value, (int, float)):
                    c.number_format = "#,##0.00;(#,##0.00)"


def main():
    if not (1 <= REPORT_MONTH <= 12):
        fail(f"REPORT_MONTH must be 1-12, got {REPORT_MONTH}.")
    period = f"{REPORT_YEAR}{REPORT_MONTH:02d}"

    sub_path = find_input("ar_subledger", period)
    gl_path = find_input("gl_trial_balance", period)

    sub = load_subledger(sub_path)
    gl_balance = load_gl_balance(gl_path, period)

    # BUSINESS RULE:
    # Intercompany customers are excluded because they are
    # reconciled separately under the intercompany close process
    # (their balances sit in GL 1210, not 1200).
    is_ic = sub["intercompany_flag"] == "Y"
    intercompany = sub[is_ic]
    trade = sub[~is_ic]

    gross_total = round(sub["open_balance"].sum(), 2)
    ic_total = round(intercompany["open_balance"].sum(), 2)
    adjusted_total = round(trade["open_balance"].sum(), 2)
    difference = round(adjusted_total - gl_balance, 2)
    ties = abs(difference) <= TOLERANCE
    result = "TIES" if ties else "DOES NOT TIE - review Exceptions tab"

    # --- Recon tab: the tie-out, then customer subtotals as support ---
    recon = pd.DataFrame([
        ("Gross AR subledger total", gross_total),
        ("Less: intercompany customers excluded", -ic_total),
        ("Adjusted AR subledger total", adjusted_total),
        (f"GL account {GL_AR_ACCOUNT} balance", gl_balance),
        ("Difference (subledger - GL)", difference),
        ("Tolerance", TOLERANCE),
        ("Result", result),
    ], columns=["Line", "Amount"])

    by_customer = (trade.groupby(["customer_id", "customer_name"], as_index=False)
                   ["open_balance"].sum().round(2))

    # --- Exceptions tab: anything a reviewer must look at ---
    exceptions = []
    if not ties:
        exceptions.append({
            "Type": "Tie-out difference",
            "Reference": f"GL {GL_AR_ACCOUNT}",
            "Customer": "",
            "Amount": difference,
            "Action needed": "Unexplained. Investigate before sign-off "
                             "(manual JEs to 1200, cash timing, missing invoices).",
        })
    for _, row in intercompany.iterrows():
        exceptions.append({
            "Type": "Excluded - intercompany (information)",
            "Reference": row["invoice_number"],
            "Customer": f"{row['customer_id']} {row['customer_name']}",
            "Amount": round(row["open_balance"], 2),
            "Action needed": "Confirm it is reconciled in the intercompany close (GL 1210).",
        })
    exceptions_df = pd.DataFrame(
        exceptions, columns=["Type", "Reference", "Customer", "Amount", "Action needed"])

    # --- Control tab: the trail from this output back to the workflow ---
    control = pd.DataFrame([
        ("Workflow name", WORKFLOW_NAME),
        ("Workflow version", WORKFLOW_VERSION),
        ("Reporting period", period),
        ("Source period (AR subledger file)", file_period(sub_path)),
        ("Source period (GL trial balance file)", file_period(gl_path)),
        ("Source files", f"{sub_path.name}; {gl_path.name}"),
        ("Records read / excluded / included",
         f"{len(sub)} / {len(intercompany)} / {len(trade)}"),
        ("Repository location", REPOSITORY_LOCATION),
        ("Git commit", git_commit()),
        ("Run date", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Run by", getpass.getuser()),
        ("Result", result),
        ("Reviewed by / date", ""),
    ], columns=["Field", "Value"])

    OUTPUT_DIR.mkdir(exist_ok=True)
    out_path = OUTPUT_DIR / f"AR_1200_Recon_{period}.xlsx"
    try:
        with pd.ExcelWriter(out_path, engine="openpyxl") as xl:
            recon.to_excel(xl, sheet_name="Recon", index=False)
            by_customer.to_excel(xl, sheet_name="Recon", index=False,
                                 startrow=len(recon) + 2)
            exceptions_df.to_excel(xl, sheet_name="Exceptions", index=False)
            control.to_excel(xl, sheet_name="Control", index=False)
            format_sheet(xl.sheets["Recon"], money_cols=(2, 3))
            format_sheet(xl.sheets["Exceptions"], money_cols=(4,))
            format_sheet(xl.sheets["Control"])
    except PermissionError:
        fail(f"Could not write {out_path.name}. Is it open in Excel? Close it and re-run.")

    print(f"{WORKFLOW_NAME} - period {period}")
    for line, amount in recon.itertuples(index=False):
        if isinstance(amount, float) and line != "Tolerance":
            print(f"  {line:<40}{amount:>12,.2f}")
    print(f"  Result: {result}")
    print(f"  Output: {out_path}")
    print("\nThe run finished. That is not the same as the recon being right. "
          "Review the workbook before filing it.")


if __name__ == "__main__":
    main()
